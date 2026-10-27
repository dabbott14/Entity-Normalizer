"""Write a normalized workbook with the same visual cues as the original tool:

* changed name cells in bold red,
* raw snapshot columns shaded grey with a "do not edit" note,
* Amount columns in accounting format,
* plus "Mapping Log" and "Alias Table" sheets for the audit trail.
"""

from __future__ import annotations

from typing import TYPE_CHECKING

import pandas as pd
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .core import ALIAS_SHEET, LOG_SHEET

if TYPE_CHECKING:
    from .core import NormalizerSession

ACCOUNTING_FORMAT = '_($* #,##0.00_);_($* (#,##0.00);_($* "-"??_);_(@_)'
CHANGED_FONT = Font(bold=True, color="FFFF0000")
RAW_FILL = PatternFill("solid", start_color="FFDCDCDC")
HEADER_FONT = Font(bold=True)


def _autofit(ws, max_width: int = 60) -> None:
    for col_cells in ws.columns:
        width = max((len(str(c.value)) for c in col_cells if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(width + 2, max_width)


def write_workbook(session: "NormalizerSession", target) -> None:
    current = session.current_sheets()

    with pd.ExcelWriter(target, engine="openpyxl") as writer:
        for name, df in current.items():
            df.to_excel(writer, sheet_name=name, index=False)
        session.mapping_log().to_excel(writer, sheet_name=LOG_SHEET, index=False)
        session.alias_table().to_excel(writer, sheet_name=ALIAS_SHEET, index=False)

        wb = writer.book
        for name, df in current.items():
            ws = wb[name]
            headers = list(df.columns)
            col_index = {h: i + 1 for i, h in enumerate(headers)}

            for cell in ws[1]:
                cell.font = HEADER_FONT
            ws.freeze_panes = "A2"

            if "Amount" in col_index:
                c = col_index["Amount"]
                for row in range(2, len(df) + 2):
                    ws.cell(row=row, column=c).number_format = ACCOUNTING_FORMAT

            for f in session.fields:
                if f.sheet != name:
                    continue
                name_col = col_index[f.column]
                raw_col = col_index[f.raw_column]

                # Highlight every cell whose value differs from the raw snapshot.
                changed = df[f.column].fillna("").astype(str) != df[f.raw_column].fillna("").astype(str)
                for offset in changed[changed].index:
                    ws.cell(row=df.index.get_loc(offset) + 2, column=name_col).font = CHANGED_FONT

                for row in range(1, len(df) + 2):
                    ws.cell(row=row, column=raw_col).fill = RAW_FILL
                ws.cell(row=1, column=raw_col).comment = Comment(
                    "Do not edit — raw snapshot of original data", "Entity Normalizer"
                )

            _autofit(ws)

            # Show date columns as dates, without a 00:00:00 time.
            for col, idx in col_index.items():
                if pd.api.types.is_datetime64_any_dtype(df[col]):
                    for row in range(2, len(df) + 2):
                        ws.cell(row=row, column=idx).number_format = "mm/dd/yyyy"
                    ws.column_dimensions[get_column_letter(idx)].width = 12

            # Width from raw numbers ("7684.13") is too narrow once accounting
            # format adds "$", commas and padding, so size from the formatted text.
            if "Amount" in col_index:
                amounts = pd.to_numeric(df["Amount"], errors="coerce").abs().dropna()
                widest = len(f"$ {amounts.max():,.2f} ") if len(amounts) else 0
                letter = get_column_letter(col_index["Amount"])
                ws.column_dimensions[letter].width = max(widest + 4, 14)

        for extra in (LOG_SHEET, ALIAS_SHEET):
            ws = wb[extra]
            for cell in ws[1]:
                cell.font = HEADER_FONT
            ws.freeze_panes = "A2"
            _autofit(ws, max_width=80)
