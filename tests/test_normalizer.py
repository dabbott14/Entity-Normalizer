"""Tests for the Entity Normalizer.

Run with either:
    python -m pytest
    python -m unittest discover tests
"""

import io
import sys
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "data"))

from normalizer import NormalizerSession, comparison_key, match_key, strip_swift_prefix, suggest_clusters  # noqa: E402
from generate_sample import build  # noqa: E402


def small_case():
    wire = pd.DataFrame(
        {
            "Orig Acct Name": ["1/ACME CORP", "Acme Corp", "SUBJECT LLC", "UNIT 21/B LLC"],
            "Bene Acct Name": ["SUBJECT LLC", "SUBJECT LLC", "acme  corp", "SUBJECT LLC"],
            "Amount": [100.0, 250.5, 75.0, 10.0],
        }
    )
    combined = pd.DataFrame(
        {
            "Tab Name": ["Wire", "Wire", "Check"],
            "Description": ["ACME CORP", "SUBJECT LLC", "ACME CORP"],
            "Account Name": ["SUBJECT LLC", "Acme Corp", "SUBJECT LLC"],
            "Amount": [100.0, 75.0, 20.0],
        }
    )
    return {"Wire": wire, "Combined": combined}


class TextCleaning(unittest.TestCase):
    def test_strips_leading_swift_prefix_only(self):
        self.assertEqual(strip_swift_prefix("1/ACME CORP"), "ACME CORP")
        self.assertEqual(strip_swift_prefix("  1/ ACME CORP "), "ACME CORP")
        # The original Replace() would have turned this into "UNIT 2B LLC"
        self.assertEqual(strip_swift_prefix("UNIT 21/B LLC"), "UNIT 21/B LLC")

    def test_non_strings_pass_through(self):
        self.assertIsNone(strip_swift_prefix(None))
        self.assertTrue(pd.isna(strip_swift_prefix(float("nan"))))

    def test_match_key_ignores_case_and_spacing(self):
        self.assertEqual(match_key("Acme  Corp "), match_key("ACME CORP"))


class Session(unittest.TestCase):
    def setUp(self):
        self.s = NormalizerSession(small_case())

    def test_raw_columns_are_appended_at_the_end(self):
        cols = list(self.s.sheets["Wire"].columns)
        self.assertEqual(cols[-2:], ["Orig Acct Name Raw", "Bene Acct Name Raw"])
        self.assertEqual(self.s.sheets["Wire"]["Orig Acct Name Raw"][0], "1/ACME CORP")

    def test_case_variants_are_one_entity(self):
        names = self.s.distinct_names().set_index("key")
        self.assertEqual(names.loc["acme corp", "Spellings"], 3)  # ACME CORP, Acme Corp, acme corp
        # Wire: 3 cells; Combined wire rows: 2 cells; Check row excluded
        self.assertEqual(names.loc["acme corp", "Count"], 5)

    def test_rename_changes_every_spelling(self):
        """Regression: the VBA version renamed only the exact spelling shown."""
        step = self.s.rename(["Acme Corp"], "Acme Corporation")
        cur = self.s.current_sheets()
        wire_names = set(cur["Wire"]["Orig Acct Name"]) | set(cur["Wire"]["Bene Acct Name"])
        self.assertNotIn("ACME CORP", wire_names)
        self.assertNotIn("acme  corp", wire_names)
        self.assertEqual(step.cells_changed, 5)

    def test_rename_scope_matches_the_list(self):
        self.s.rename(["ACME CORP"], "Acme Corporation")
        combined = self.s.current_sheets()["Combined"]
        self.assertEqual(combined.loc[2, "Description"], "ACME CORP")  # Check row untouched

    def test_rename_all_rows_option(self):
        s = NormalizerSession(small_case(), rename_all_rows=True)
        s.rename(["ACME CORP"], "Acme Corporation")
        self.assertEqual(s.current_sheets()["Combined"].loc[2, "Description"], "Acme Corporation")

    def test_multi_level_undo(self):
        self.s.rename(["ACME CORP"], "Acme Corporation")
        self.s.rename(["Acme Corporation"], "ACME CORPORATION")
        self.s.undo()
        self.assertIn("Acme Corporation", set(self.s.current_sheets()["Wire"]["Orig Acct Name"]))
        self.s.undo()
        self.assertEqual(self.s.current_sheets()["Wire"].loc[0, "Orig Acct Name"], "ACME CORP")
        self.assertIsNone(self.s.undo())

    def test_raw_columns_never_change(self):
        self.s.rename(["ACME CORP"], "Acme Corporation")
        self.assertEqual(self.s.current_sheets()["Wire"].loc[0, "Orig Acct Name Raw"], "1/ACME CORP")

    def test_validation(self):
        with self.assertRaises(ValueError):
            self.s.rename(["ACME CORP"], "   ")
        with self.assertRaises(ValueError):
            self.s.rename([], "X")

    def test_suggested_name_is_most_common_spelling(self):
        self.assertEqual(self.s.suggested_name(["acme corp"]), "ACME CORP")

    def test_alias_table_round_trip(self):
        self.s.rename(["ACME CORP"], "Acme Corporation")
        aliases = self.s.alias_table()
        fresh = NormalizerSession(small_case())
        self.assertEqual(fresh.apply_alias_table(aliases), 1)
        self.assertTrue(fresh.current_sheets()["Wire"].equals(self.s.current_sheets()["Wire"]))

    def test_missing_sheet_or_column_warns(self):
        s = NormalizerSession({"Wire": pd.DataFrame({"Orig Acct Name": ["A"]})})
        self.assertTrue(any("Bene Acct Name" in w for w in s.warnings))
        self.assertTrue(any("Combined" in w for w in s.warnings))


class Export(unittest.TestCase):
    def setUp(self):
        self.s = NormalizerSession(small_case())
        self.s.rename(["ACME CORP"], "Acme Corporation")
        self.buffer = io.BytesIO(self.s.to_excel_bytes())

    def test_formatting(self):
        wb = load_workbook(self.buffer)
        ws = wb["Wire"]
        headers = [c.value for c in ws[1]]
        name_col = headers.index("Orig Acct Name") + 1
        raw_col = headers.index("Orig Acct Name Raw") + 1
        amount_col = headers.index("Amount") + 1
        self.assertTrue(ws.cell(2, name_col).font.bold)          # renamed -> red bold
        self.assertFalse(ws.cell(4, name_col).font.bold)         # SUBJECT LLC unchanged
        self.assertIsNotNone(ws.cell(1, raw_col).comment)
        self.assertEqual(ws.cell(2, raw_col).fill.start_color.rgb, "FFDCDCDC")
        self.assertIn("#,##0.00", ws.cell(2, amount_col).number_format)
        self.assertIn("Mapping Log", wb.sheetnames)
        self.assertIn("Alias Table", wb.sheetnames)

    def test_reopening_export_restores_work(self):
        reopened = NormalizerSession.from_excel(self.buffer)
        self.assertEqual(reopened.restored_steps, 1)
        self.assertNotIn("Mapping Log", reopened.sheets)
        self.assertEqual(
            reopened.current_sheets()["Wire"].loc[0, "Orig Acct Name"], "Acme Corporation"
        )
        # Raw snapshot survives the round trip unchanged
        self.assertEqual(reopened.sheets["Wire"].loc[0, "Orig Acct Name Raw"], "1/ACME CORP")


class Suggestions(unittest.TestCase):
    def test_comparison_key(self):
        self.assertEqual(comparison_key("Meridian Coastal Bk, N.A."), "MERIDIAN COASTAL BANK")
        self.assertEqual(comparison_key("J P Morgan"), "JP MORGAN")
        self.assertEqual(comparison_key("1/Silverline Hldgs Inc"), "SILVERLINE HOLDINGS")

    def test_groups_variants_and_word_order(self):
        clusters = suggest_clusters(
            ["HARBOR POINT LOGISTICS LLC", "Harbor Point Logistics, LLC", "DELANEY, MARCUS T",
             "MARCUS T DELANEY", "BLUE WATER TRADING CO", "BLUEWATER TRADING CO", "Granite Peak"],
            [5, 2, 1, 4, 1, 3, 9],
        )
        groups = [set(c["names"]) for c in clusters]
        self.assertIn({"HARBOR POINT LOGISTICS LLC", "Harbor Point Logistics, LLC"}, groups)
        self.assertIn({"DELANEY, MARCUS T", "MARCUS T DELANEY"}, groups)
        self.assertIn({"BLUE WATER TRADING CO", "BLUEWATER TRADING CO"}, groups)
        self.assertFalse(any("Granite Peak" in g for g in groups))
        self.assertEqual(clusters[0]["total"], 7)  # largest first

    def test_unrelated_names_stay_apart(self):
        self.assertEqual(suggest_clusters(["NORTHGATE SERVICES", "SOUTHPORT TRADING"], [1, 1]), [])

    def test_sample_data_end_to_end(self):
        s = NormalizerSession(build())
        names = s.distinct_names()
        clusters = suggest_clusters(names["Name"], names["Count"])
        self.assertGreaterEqual(len(clusters), 10)
        before = len(names)
        for c in clusters:
            s.rename(c["names"], c["suggested"])
        self.assertLess(len(s.distinct_names()), before)


if __name__ == "__main__":
    unittest.main()
